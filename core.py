"""محرك الستكرات: قراءة الإكسل + رسم الملصق (نص عربي/إنجليزي + QR)."""
import os
import sys
from dataclasses import dataclass
from functools import lru_cache

import arabic_reshaper
import qrcode
from bidi.algorithm import get_display
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


# ---------- الإعدادات ----------
@dataclass
class LabelSettings:
    width_mm: float = 50.0
    height_mm: float = 30.0
    margin_mm: float = 2.0
    font_pt: float = 12.0
    qr_mm: float = 20.0
    qr_pos: str = "left"      # left / right / top
    align: str = "center"     # center / left / right
    bold_first: bool = True
    border: bool = False


def mm_to_px(v, dpi):
    return int(round(float(v) / 25.4 * dpi))


def pt_to_px(v, dpi):
    return max(1, int(round(float(v) / 72 * dpi)))


# ---------- الإكسل ----------
def _fmt(c):
    if c is None:
        return ""
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c).strip()


def read_excel(path):
    """يرجع (العناوين، الصفوف). الصف الأول = العناوين."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    headers = []
    for i, h in enumerate(rows[0]):
        headers.append(_fmt(h) or f"عمود {i + 1}")
    n = len(headers)
    data = []
    for r in rows[1:]:
        r = list(r)[:n] + [None] * max(0, n - len(r))
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        data.append([_fmt(c) for c in r])
    return headers, data


# ---------- الخطوط والنص ----------
_FONT_CANDIDATES = {
    False: [
        r"C:\Windows\Fonts\tahoma.ttf", r"C:\Windows\Fonts\arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ],
    True: [
        r"C:\Windows\Fonts\tahomabd.ttf", r"C:\Windows\Fonts\arialbd.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
    ],
}


@lru_cache(maxsize=256)
def load_font(size_px, bold):
    # خط مرفق بجانب البرنامج له الأولوية (font.ttf / font-bold.ttf)
    base = os.path.dirname(os.path.abspath(sys.argv[0] if getattr(sys, "frozen", False) else __file__))
    local = os.path.join(base, "font-bold.ttf" if bold else "font.ttf")
    for p in [local] + _FONT_CANDIDATES[bold]:
        if os.path.exists(p):
            if HAS_RAQM:
                return ImageFont.truetype(p, size_px)
            return ImageFont.truetype(p, size_px, layout_engine=ImageFont.Layout.BASIC)
    return ImageFont.load_default()


try:
    from PIL import features
    HAS_RAQM = features.check("raqm")
except Exception:
    HAS_RAQM = False


def shape(text):
    # إن كان Pillow يملك raqm فهو يعالج الاتجاه والتشكيل بنفسه
    if HAS_RAQM:
        return text
    """تشكيل الحروف العربية واتجاه النص ليظهر صحيحاً في الصورة."""
    if any("\u0600" <= ch <= "\u06FF" for ch in text):
        return get_display(arabic_reshaper.reshape(text))
    return text


def _width(draw, text, font):
    return draw.textlength(shape(text), font=font) if text else 0


def wrap(draw, text, font, max_w):
    """تقسيم النص على أسطر حسب العرض المتاح."""
    lines, cur = [], ""
    for word in text.split():
        trial = (cur + " " + word).strip()
        if _width(draw, trial, font) <= max_w:
            cur = trial
            continue
        if cur:
            lines.append(cur)
            cur = ""
        # كلمة أطول من العرض: نكسرها حرفاً حرفاً
        if _width(draw, word, font) > max_w:
            piece = ""
            for ch in word:
                if _width(draw, piece + ch, font) > max_w and piece:
                    lines.append(piece)
                    piece = ch
                else:
                    piece += ch
            cur = piece
        else:
            cur = word
    if cur:
        lines.append(cur)
    return lines or [""]


def _layout_text(draw, text_lines, box_w, box_h, s, dpi):
    """يجرّب أحجام الخط من الأكبر للأصغر حتى يتسع النص داخل الصندوق."""
    size = float(s.font_pt)
    while True:
        px = pt_to_px(size, dpi)
        blocks, total = [], 0
        for i, t in enumerate(text_lines):
            f = load_font(px, bool(s.bold_first and i == 0))
            for ln in wrap(draw, t, f, box_w):
                lh = int(px * 1.25)
                blocks.append((ln, f, lh))
                total += lh
        if total <= box_h or size <= 5:
            return blocks, total
        size -= 0.5


# ---------- QR ----------
def draw_qr(img, data, area):
    x0, y0, x1, y1 = area
    side = min(x1 - x0, y1 - y0)
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, border=0)
    qr.add_data(data)
    qr.make(fit=True)
    m = qr.get_matrix()
    n = len(m)
    mod = max(1, side // n)         # وحدات صحيحة = مسح أوضح
    size = mod * n
    ox = x0 + (x1 - x0 - size) // 2
    oy = y0 + (y1 - y0 - size) // 2
    d = ImageDraw.Draw(img)
    for r, row in enumerate(m):
        for c, v in enumerate(row):
            if v:
                d.rectangle([ox + c * mod, oy + r * mod, ox + (c + 1) * mod - 1, oy + (r + 1) * mod - 1], fill=0)


# ---------- رسم الملصق ----------
def render_label(text_lines, qr_text, s, dpi=203):
    """يرجع صورة بالأبيض والأسود (mode L) بحجم الملصق الحقيقي عند الـ dpi المعطى."""
    W, H = mm_to_px(s.width_mm, dpi), mm_to_px(s.height_mm, dpi)
    img = Image.new("L", (W, H), 255)
    draw = ImageDraw.Draw(img)
    m, gap = mm_to_px(s.margin_mm, dpi), mm_to_px(1.5, dpi)
    x0, y0, x1, y1 = m, m, W - m, H - m
    iw, ih = x1 - x0, y1 - y0
    tb = (x0, y0, x1, y1)

    if qr_text:
        want = mm_to_px(s.qr_mm, dpi)
        if s.qr_pos == "top":
            qs = max(10, min(want, iw, ih - mm_to_px(6, dpi)))
            qx = x0 + (iw - qs) // 2
            draw_qr(img, qr_text, (qx, y0, qx + qs, y0 + qs))
            tb = (x0, y0 + qs + gap, x1, y1)
        else:
            qs = max(10, min(want, ih, iw - mm_to_px(10, dpi)))
            qy = y0 + (ih - qs) // 2
            if s.qr_pos == "right":
                draw_qr(img, qr_text, (x1 - qs, qy, x1, qy + qs))
                tb = (x0, y0, x1 - qs - gap, y1)
            else:
                draw_qr(img, qr_text, (x0, qy, x0 + qs, qy + qs))
                tb = (x0 + qs + gap, y0, x1, y1)

    lines = [t for t in text_lines if t]
    if lines:
        bx0, by0, bx1, by1 = tb
        bw, bh = max(1, bx1 - bx0), max(1, by1 - by0)
        blocks, total = _layout_text(draw, lines, bw, bh, s, dpi)
        y = by0 + max(0, (bh - total) // 2)
        for ln, f, lh in blocks:
            w = _width(draw, ln, f)
            if s.align == "left":
                x = bx0
            elif s.align == "right":
                x = bx1 - w
            else:
                x = bx0 + (bw - w) / 2
            draw.text((x, y), shape(ln), font=f, fill=0)
            y += lh

    if s.border:
        draw.rectangle([0, 0, W - 1, H - 1], outline=0, width=max(1, dpi // 100))

    return img.point(lambda p: 255 if p > 150 else 0)
